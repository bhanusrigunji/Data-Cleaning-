"""Excel and chart report generation."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _quality_table(data: pd.DataFrame, rules: dict[str, str]) -> pd.DataFrame:
    total = max(len(data), 1)
    return pd.DataFrame(
        {
            "Column": data.columns,
            "Data Type": [str(data[column].dtype) for column in data.columns],
            "Missing Values": [int(data[column].isna().sum()) for column in data.columns],
            "Missing %": [round(data[column].isna().sum() / total * 100, 2) for column in data.columns],
            "Fill Rule": [rules.get(column, "") for column in data.columns],
        }
    )


def write_excel_report(data: pd.DataFrame, summary: dict[str, Any], output_path: Path) -> None:
    """Write cleaned data, run summary, and per-column quality details."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    quality = _quality_table(data, summary["fill_rules"])
    summary_rows = [
        ("Rows before cleaning", summary["rows_before"]),
        ("Rows after cleaning", summary["rows_after"]),
        ("Duplicate rows removed", summary["duplicates_removed"]),
        ("Missing cells before fill", summary["missing_cells_before_fill"]),
        ("Missing cells after fill", summary["missing_cells_after_fill"]),
        ("Columns", summary["columns_after"]),
    ]
    summary_df = pd.DataFrame(summary_rows, columns=["Metric", "Value"])

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        data.to_excel(writer, index=False, sheet_name="Cleaned Data")
        summary_df.to_excel(writer, index=False, sheet_name="Summary")
        quality.to_excel(writer, index=False, sheet_name="Data Quality")

        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="234E70")
                cell.alignment = Alignment(horizontal="center")
            for cells in worksheet.columns:
                width = min(max(max(len(str(cell.value or "")) for cell in cells) + 2, 12), 42)
                worksheet.column_dimensions[get_column_letter(cells[0].column)].width = width


def write_summary_chart(summary: dict[str, Any], output_path: Path) -> None:
    """Create a compact before/after row-count chart."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    labels = ["Before cleaning", "After cleaning"]
    values = [summary["rows_before"], summary["rows_after"]]
    fig, ax = plt.subplots(figsize=(7, 4))
    bars = ax.bar(labels, values, color=["#9BB7D4", "#2A6F97"], width=0.55)
    ax.set_title("Rows in the cleaning workflow")
    ax.set_ylabel("Number of rows")
    ax.spines[["top", "right"]].set_visible(False)
    ax.bar_label(bars, padding=4, fmt="%d")
    fig.tight_layout()
    fig.savefig(output_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
